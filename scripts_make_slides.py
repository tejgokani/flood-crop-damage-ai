"""Build the Tech-a-thon presentation deck from the live results."""
from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

NAVY, SLATE, ACCENT = RGBColor(0x0F, 0x23, 0x3E), RGBColor(0x47, 0x55, 0x69), RGBColor(0x1D, 0x4E, 0xD8)
GREEN, AMBER, WHITE = RGBColor(0x15, 0x80, 0x3D), RGBColor(0xB4, 0x53, 0x09), RGBColor(0xFF, 0xFF, 0xFF)

ROOT = Path(__file__).parent
REPORTS = ROOT / "reports"


def load(name):
    p = REPORTS / name
    return json.loads(p.read_text()) if p.exists() else None


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide(title: str, subtitle: str = ""):
    s = prs.slides.add_slide(BLANK)
    tb = s.shapes.add_textbox(Inches(0.55), Inches(0.34), Inches(12.2), Inches(0.85))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.size, p.font.bold, p.font.color.rgb = Pt(30), True, NAVY
    if subtitle:
        p2 = tb.text_frame.add_paragraph()
        p2.text = subtitle
        p2.font.size, p2.font.color.rgb = Pt(13), SLATE
    return s


def bullets(s, items, left=0.7, top=1.55, width=12.0, size=15, gap=8):
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.3))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        txt, lvl = (it if isinstance(it, tuple) else (it, 0))
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("• " if lvl == 0 else "– ") + txt
        p.level = lvl
        p.font.size = Pt(size if lvl == 0 else size - 2)
        p.font.color.rgb = NAVY if lvl == 0 else SLATE
        p.space_after = Pt(gap)
    return tb


def table(s, rows, left, top, width, height, header_fill=NAVY, fs=10.5):
    r, c = len(rows), len(rows[0])
    shape = s.shapes.add_table(r, c, Inches(left), Inches(top), Inches(width), Inches(height))
    t = shape.table
    for j, val in enumerate(rows[0]):
        cell = t.cell(0, j)
        cell.text = str(val)
        cell.fill.solid(); cell.fill.fore_color.rgb = header_fill
        para = cell.text_frame.paragraphs[0]
        para.font.size, para.font.bold, para.font.color.rgb = Pt(fs), True, WHITE
    for i in range(1, r):
        for j in range(c):
            cell = t.cell(i, j)
            cell.text = str(rows[i][j])
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(fs)
            para.font.color.rgb = NAVY
    return t


# ---------------------------------------------------------------- 1 title
s = prs.slides.add_slide(BLANK)
box = s.shapes.add_textbox(Inches(0.9), Inches(2.1), Inches(11.5), Inches(2.6))
tf = box.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "AI-Based Crop Flood Damage Assessment"
p.font.size, p.font.bold, p.font.color.rgb = Pt(42), True, NAVY
for txt, sz, col in [
    ("Agriculture + Disaster Management  ·  Problem Statement 6", 18, ACCENT),
    ("Five hybrid deep learning architectures benchmarked on one protocol,", 15, SLATE),
    ("fused with Indian district crop statistics", 15, SLATE),
    ("", 8, SLATE),
    ("VIT Tech-a-thon · Fall Semester 2026–27", 13, SLATE),
    ("github.com/tejgokani/flood-crop-damage-ai", 12, ACCENT),
]:
    q = tf.add_paragraph(); q.text = txt
    q.font.size, q.font.color.rgb = Pt(sz), col

# ---------------------------------------------------------------- 2 problem
s = slide("The problem", "Why field surveys fail exactly when they are needed")
bullets(s, [
    "Floods destroy crops across large geographical regions.",
    "Traditional field surveys are impossible immediately after a disaster — roads and agricultural land are inaccessible.",
    "Relief and insurance decisions still have to be made, within days.",
    "",
    ("Required output: Healthy → Mild → Moderate → Severe crop damage", 0),
    ("Required approach: satellite / drone imagery + deep learning, five hybrid techniques", 0),
], size=16, gap=12)

# ---------------------------------------------------------------- 3 objective from papers
s = slide("Objective — derived from six papers' own future work",
          "Every quote below is verbatim from the source paper's Future Work / Limitations section")
rows = [["Paper (2026)", "Its stated future work — verbatim", "What we built"],
        ["FLNet\narXiv:2601.03884", "\"…experimenting with state-of-the-art backbones like Transformers\"; \"handle cloud cover\"", "Swin + 5-model benchmark;\nSAR instead of optical"],
        ["Swin flood scene\nSci Rep 2026", "\"combining temporal modeling…\"; \"future research will concentrate on incorporating multi-spectral or hydrological inputs\"", "CNN+LSTM;\nSentinel-1 SAR not RGB"],
        ["Physics-guided U-Net+FNO\nRadarConf 2026", "\"relaxing flow assumptions, optimizing memory usage, and enhancing regularization\"", "Automatic over/under-fit\ncorrection; capped compute"],
        ["Explainable SAR CNN/Transformer\narXiv:2606.16302", "\"incorporating multi-temporal SAR data, in which pre-flood and post-flood images are jointly analyzed\"", "Real pre/post pairs;\npermanent water subtracted"],
        ["TDAVM-UNet\nFrontiers Plant Sci 2026", "\"Future work: dataset expansion, domain adaptation, edge deployment…\"", "GAN augmentation;\nprogressive data tiering"],
        ["Climate-informed cropland\nSci Rep", "\"…can be used to study climate change impacts on the data-scarce, critical flood zones\"", "Indian district crop fusion\n(a data-scarce flood zone)"]]
t = table(s, rows, 0.5, 1.5, 12.35, 5.2, fs=9)
t.columns[0].width, t.columns[1].width, t.columns[2].width = Emu(Inches(2.9)), Emu(Inches(6.1)), Emu(Inches(3.35))

# ---------------------------------------------------------------- 4 objective statement
s = slide("The objective statement")
tb = s.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.8), Inches(2.0))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]
p.text = ("To design and evaluate a cloud-penetrating, multi-temporal, multi-architecture deep "
          "learning system that identifies agricultural fields from Sentinel-1 SAR imagery, "
          "classifies flood-induced crop damage as Healthy → Mild → Moderate → Severe, and fuses "
          "that severity with district-level Indian agricultural statistics to produce a "
          "region-scale crop loss assessment usable when field surveys are impossible.")
p.font.size, p.font.color.rgb, p.font.italic = Pt(17), NAVY, True
bullets(s, [
    "SO-1  Benchmark five hybrids under one identical protocol            → gap G3",
    "SO-2  Overcome data scarcity without touching evaluation integrity   → gap G1",
    "SO-3  Model flood evolution over time, not a single snapshot         → gap G2",
    "SO-4  Fuse imagery with real regional agricultural context           → gap G4",
    "SO-5  Make regularisation automatic and keep compute feasible        → gap G6",
], top=3.7, size=14, gap=9)

# ---------------------------------------------------------------- 5 flowchart
s = slide("Flow chart of the proposed work", "Both prescribed sequences, implemented in the given order")
fc = ROOT / "docs" / "flowchart.png"
if fc.exists():
    s.shapes.add_picture(str(fc), Inches(0.85), Inches(1.25), height=Inches(5.95))

# ---------------------------------------------------------------- 6 sequence compliance
s = slide("Compliance with the prescribed sequence", "Every step maps to a function that actually runs")
rows = [["#", "Step", "Implemented in"],
        ["1", "Load dataset", "data/etci.py · data/india_agri.py"],
        ["2", "Data preprocessing", "preprocess/transforms.py"],
        ["3", "Data leakage checks — target · duplicate · suspicious", "preprocess/leakage.py"],
        ["4", "CLAHE", "preprocess/clahe.py"],
        ["5", "Split / 10-fold cross validation", "splits.py"],
        ["6", "GAN augmentation / SMOTE — TRAINING DATA ONLY", "augment/gan.py · augment/smote.py"],
        ["7", "Initial model training", "train/loop.py"],
        ["8", "Overfitting / underfitting detection", "train/diagnostics.py"],
        ["9", "Apply correction technique and retrain", "train/correction.py"],
        ["10", "Final test evaluation", "eval/metrics.py"]]
t = table(s, rows, 0.6, 1.5, 12.1, 4.6, fs=11)
t.columns[0].width, t.columns[1].width, t.columns[2].width = Emu(Inches(0.5)), Emu(Inches(6.6)), Emu(Inches(5.0))
tb = s.shapes.add_textbox(Inches(0.6), Inches(6.25), Inches(12.1), Inches(0.8))
p = tb.text_frame.paragraphs[0]
p.text = ("Why augmentation comes after the split: SMOTE interpolates between neighbours, so fitting it "
          "before the split lets a test row be interpolated into a training row — the score then measures "
          "memorisation. tests/test_no_leakage.py asserts val/test are never resampled.")
p.font.size, p.font.color.rgb, p.font.italic = Pt(11), AMBER, True
tb.text_frame.word_wrap = True

# ---------------------------------------------------------------- 7 architectures
s = slide("The five hybrid architectures", "Shared U-Net decoder and shared severity head — so the comparison isolates the encoder")
rows = [["#", "Hybrid", "Encoder", "Why it is in the benchmark"],
        ["1", "YOLO12 + U-Net", "R-ELAN + area attention\n(implemented natively)", "Detects field parcels while segmenting water inside them"],
        ["2", "ResNet + U-Net", "timm resnet34", "The established baseline to measure against"],
        ["3", "EfficientNet + Attention", "efficientnet_b0 + CBAM\n+ attention-gated skips", "Best accuracy per FLOP; the shape TDAVM-UNet validates for agriculture"],
        ["4", "Swin Transformer + U-Net", "swinv2_tiny", "Global receptive field at linear cost — long-range boundary precision"],
        ["5", "CNN + LSTM", "CNN + ConvLSTM at every scale", "Consumes the real pre-monsoon / monsoon pair"]]
t = table(s, rows, 0.5, 1.5, 12.35, 4.3, fs=10)
for col, w in zip(t.columns, [0.45, 2.8, 3.1, 6.0]):
    col.width = Emu(Inches(w))
tb = s.shapes.add_textbox(Inches(0.5), Inches(6.0), Inches(12.35), Inches(1.1))
tb.text_frame.word_wrap = True
p = tb.text_frame.paragraphs[0]
p.text = ("Shared head: the severity classifier reads average-pooled features, max-pooled features, and "
          "mean(sigmoid(segmentation)) — the predicted net flood fraction. Since the severity label is "
          "defined as that fraction, the classifier is handed the exact quantity the target is built from.")
p.font.size, p.font.color.rgb = Pt(11), SLATE

# ---------------------------------------------------------------- 8 data
s = slide("Data", "Both sources public, no credentials, nothing bulk-downloaded")
bullets(s, [
    "Imagery — ETCI-2021 Sentinel-1 SAR over Bangladesh (Ganges–Brahmaputra delta)",
    ("2,032 tiles with a genuine pre/post pair: 2017-03-14 pre-monsoon and 2017-06-06 monsoon flood", 1),
    ("Each carries VV, VH, a flood mask and a permanent-water mask at both dates", 1),
    "Agriculture — ICRISAT district statistics: 310 Indian districts · 23 crops · 2010–2017",
    ("Label from per-district yield anomaly → a real 10.9 : 1 class imbalance, which is what justifies SMOTE", 1),
    "",
    "Severity = net flood fraction (flood − permanent water), binned at 2% / 10% / 33%",
    ("The 33% boundary is the Indian NDRF/SDRF crop-loss relief threshold — not a round number", 1),
    ("Permanent water is subtracted because a tile containing a river is not a damaged tile", 1),
], size=14, gap=8)

# ---------------------------------------------------------------- 9 progressive scaling
s = slide("Progressive data scaling", "Start small; step up only as far as the machine demonstrably sustains")
rows = [["Tier", "Tiles", "Size", "Purpose"],
        ["T0_smoke", "40", "128 px", "Pipeline proof — offline, runs in CI"],
        ["T1_tiny", "150", "128 px", "First real data"],
        ["T2_small", "400", "192 px", "First meaningful signal"],
        ["T3_medium", "900", "256 px", "Target tier"],
        ["T4_large", "1800", "256 px", "Only if T3 finished comfortably"]]
table(s, rows, 0.7, 1.5, 7.2, 2.9, fs=11)
bullets(s, [
    "After each tier we measure peak RSS, free disk and seconds/epoch",
    "Advance only if the next tier stays under a 6 GB disk floor, 70% of RAM and the time budget",
    "Otherwise stop at the last good tier and record why — in reports/scaling_log.json",
    "A complete five-model table exists after the first tier, so there is always something to show",
], left=8.2, top=1.6, width=4.6, size=12, gap=10)

# ---------------------------------------------------------------- 10 results
s = slide("Results", "")
results, scaling, tabular = load("results.json"), load("scaling_log.json"), load("tabular.json")
if results and results.get("tiers"):
    from importlib import import_module
    import sys
    sys.path.insert(0, str(ROOT / "src"))
    DISPLAY = import_module("fcda.models.registry").DISPLAY_NAMES
    tier = results["tiers"][-1]
    s.shapes.title  # noqa
    hdr = [["Hybrid", "Params", "Test macro-F1", "Accuracy", "κ", "Flood IoU", "Diagnosis"]]
    ranked = sorted(tier["models"], key=lambda m: (m.get("final_test") or {}).get("macro_f1", -1), reverse=True)
    for m in ranked:
        t_ = m.get("final_test") or {}
        hdr.append([DISPLAY.get(m["name"], m["name"]),
                    f"{(m.get('initial') or {}).get('n_params', 0)/1e6:.1f}M",
                    f"{t_.get('macro_f1', 0):.3f}", f"{t_.get('accuracy', 0):.3f}",
                    f"{t_.get('kappa', 0):.3f}", f"{t_.get('iou', 0):.3f}",
                    (m.get("diagnosis") or {}).get("status", "—")])
    table(s, hdr, 0.5, 1.45, 12.35, 2.6, fs=10.5)
    sub = s.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(12.35), Inches(2.6))
    sub.text_frame.word_wrap = True
    p = sub.text_frame.paragraphs[0]
    p.text = (f"Deepest tier completed: {tier['tier']} — {tier['n_tiles']} tiles at "
              f"{tier['image_size']}px on {tier['device']}.  Class distribution "
              f"{tier.get('class_distribution', {})}.")
    p.font.size, p.font.color.rgb = Pt(12), SLATE
    if tabular:
        tm, cv = tabular.get("test_metrics", {}), tabular.get("cv", {})
        q = sub.text_frame.add_paragraph()
        q.text = (f"Tabular pipeline — test macro-F1 {tm.get('macro_f1', 0):.3f}, "
                  f"10-fold CV {cv.get('mean_macro_f1', 0):.3f} ± {cv.get('std', 0):.3f}. "
                  f"Target leakage caught and dropped: {', '.join(tabular.get('dropped_columns', []))}.")
        q.font.size, q.font.color.rgb = Pt(12), SLATE
    if scaling:
        q = sub.text_frame.add_paragraph()
        q.text = f"Scaling stopped because: {scaling.get('stopped_because', '')}"
        q.font.size, q.font.color.rgb, q.font.italic = Pt(11), AMBER, True
else:
    bullets(s, ["Training run in progress — regenerate this deck with `python scripts_make_slides.py`."])

# ---------------------------------------------------------------- 11 figures
for fig, title in (("learning_curves.png", "Learning curves — initial training and after correction"),
                   ("confusion_matrices.png", "Confusion matrices on the held-out test split"),
                   ("scaling.png", "What the machine actually sustained")):
    fp = REPORTS / "figures" / fig
    if fp.exists():
        s = slide(title)
        s.shapes.add_picture(str(fp), Inches(0.7), Inches(1.6), width=Inches(11.9))

# ---------------------------------------------------------------- 12 demo
s = slide("Working demo", "streamlit run app/streamlit_app.py")
bullets(s, [
    "Select or upload a SAR tile → predicted damage mask overlaid on the backscatter image",
    "Severity with confidence, net inundated fraction, permanent water shown separately",
    "District loss estimate: affected hectares, production loss in tonnes, indicative value in ₹ crore",
    "Fusion panel showing whether imagery and district history agree, and whether severity was escalated",
    "",
    "Runs with or without a trained checkpoint — the rule-based path still demonstrates preprocessing, "
    "the severity rule and the fusion arithmetic",
], size=14, gap=10)

# ---------------------------------------------------------------- 13 limitations
s = slide("Honest limitations", "Stated because they are the questions worth asking")
bullets(s, [
    "Severity is derived from flood extent, not measured agronomic damage — ETCI-2021 has flood masks, not crop-damage ground truth. Validation against field-surveyed loss is NOT claimed.",
    "Only 20 Severe tiles exist in the whole pool (0.98% natural prior), so Severe-class F1 is noisy by construction. This is the corpus ceiling and is what GAN augmentation pushes against.",
    "Tiers cap Healthy at 45% to make the benchmark usable; the natural 86.9% prior is reported alongside every result, not hidden.",
    "Training geography is the Ganges–Brahmaputra delta — a reasonable analogue for eastern India, but cross-region transfer is untested. FLNet states the same limitation about itself.",
    "Rupee figures use assumed per-severity loss coefficients, surfaced as parameters rather than fitted values.",
], size=13, gap=13)

# ---------------------------------------------------------------- 14 remaining
s = slide("What remains — the other 40%")
bullets(s, [
    "Validation against field-surveyed crop loss (BFCD-22, or state revenue records)",
    "Cross-region transfer to Indian districts, and domain adaptation",
    "Sentinel-1 + Sentinel-2 fusion, as FLNet's future work proposes",
    "Calibration and uncertainty estimates, as the Explainable SAR paper proposes",
    "10-fold cross validation extended to the image pipeline once compute allows",
    "Higher tiers (T4) and longer schedules on stronger hardware",
], size=15, gap=13)

# ---------------------------------------------------------------- 15 close
s = slide("Summary")
bullets(s, [
    "Objective derived from six papers' own stated future work — quoted verbatim, not paraphrased",
    "Both prescribed sequences implemented in full, every step mapped to a function",
    "Five hybrid architectures on one shared protocol — the comparison isolates the encoder",
    "Leakage checks that produce real findings; augmentation confined to the training split by construction",
    "Automatic over/under-fit detection and correction, with before/after recorded",
    "Fusion turns predicted pixels into hectares, tonnes and rupees for a real Indian district",
    "",
    "github.com/tejgokani/flood-crop-damage-ai",
], size=14, gap=11)

out = ROOT / "slides" / "Techathon_PS6.pptx"
out.parent.mkdir(exist_ok=True)
prs.save(out)
print(f"wrote {out} — {len(prs.slides.__iter__.__self__._sldIdLst)} slides")
